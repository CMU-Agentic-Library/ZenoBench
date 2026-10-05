"""Build an editable 12-slide ZenoBench group-meeting deck.

Run from the repository root with python-pptx installed.
"""

from __future__ import annotations

from pathlib import Path
import json
import cv2

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentations"
OUT.mkdir(exist_ok=True)
FRAMES = OUT / "frames"
FRAMES.mkdir(exist_ok=True)
DEST = OUT / "ZenoBench_group_meeting.pptx"

PPT = Presentation()
PPT.slide_width = Inches(13.333)
PPT.slide_height = Inches(7.5)
BLANK = PPT.slide_layouts[6]

NAVY = RGBColor(16, 31, 52)
INK = RGBColor(25, 40, 61)
BLUE = RGBColor(36, 104, 188)
TEAL = RGBColor(0, 143, 147)
AMBER = RGBColor(239, 163, 60)
PALE = RGBColor(241, 246, 251)
MID = RGBColor(88, 105, 123)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(38, 147, 101)
RED = RGBColor(196, 78, 75)
LINE = RGBColor(219, 228, 236)


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background() if line is None else None
    if line is not None:
        s.line.color.rgb = line
    return s


def text(slide, content, x, y, w, h, size=18, color=INK, bold=False,
         font="Aptos", align=None, margin=0, valign=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(str(content).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
        p.space_after = Pt(3)
        if align is not None: p.alignment = align
    return s


def pic(slide, rel, x, y, w, h, crop=True):
    path = ROOT / rel
    from PIL import Image
    im = Image.open(path)
    iw, ih = im.size
    if not crop:
        scale = min(w / iw, h / ih)
        ww, hh = iw * scale, ih * scale
        return slide.shapes.add_picture(str(path), Inches(x + (w - ww) / 2),
                                        Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))
    ratio = w / h
    ir = iw / ih
    if ir > ratio:
        side = (1 - ratio / ir) / 2
        s = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        s.crop_left = s.crop_right = side
    else:
        side = (1 - ir / ratio) / 2
        s = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        s.crop_top = s.crop_bottom = side
    return s


def slide(title, kicker=None, subtitle=None, dark=False):
    s = PPT.slides.add_slide(BLANK)
    box(s, 0, 0, 13.333, 7.5, NAVY if dark else WHITE)
    fg = WHITE if dark else INK
    if kicker: text(s, kicker.upper(), .55, .28, 10.8, .23, 10, TEAL if not dark else AMBER, True)
    text(s, title, .55, .64, 12.05, .75, 28, fg, True)
    if subtitle: text(s, subtitle, .57, 1.38, 12.0, .43, 12.5, MID if not dark else RGBColor(181, 200, 219))
    box(s, .55, 7.08, 12.23, .012, LINE if not dark else RGBColor(81, 102, 124))
    text(s, "ZenoBench  ·  CMU Agentic Library", .56, 7.13, 8, .18, 8.5,
         MID if not dark else RGBColor(180, 196, 215))
    text(s, f"{len(PPT.slides):02d} / 12", 12.14, 7.12, .62, .2, 8.5,
         MID if not dark else RGBColor(180, 196, 215), align=PP_ALIGN.RIGHT)
    return s


def card(slide, x, y, w, h, head, body, accent=BLUE, bg=PALE, head_size=17, body_size=13):
    box(slide, x, y, w, h, bg, LINE, radius=True)
    box(slide, x, y, .055, h, accent)
    if h < 1.6:
        text(slide, head, x + .22, y + .12, w - .45, .33, head_size, INK, True)
        text(slide, body, x + .22, y + .48, w - .45, h - .54, body_size, MID)
    else:
        text(slide, head, x + .22, y + .2, w - .45, .42, head_size, INK, True)
        text(slide, body, x + .22, y + .72, w - .45, h - .85, body_size, MID)


def note(s, txt):
    try:
        s.notes_slide.notes_text_frame.text = txt
    except Exception:
        pass


def frame(video, frac, name):
    dst = FRAMES / f"{name}.jpg"
    if dst.exists(): return str(dst.relative_to(ROOT))
    cap = cv2.VideoCapture(str(ROOT / "media/tasks" / video))
    total = cap.get(cv2.CAP_PROP_FRAME_COUNT)
    cap.set(cv2.CAP_PROP_POS_FRAMES, int(total * frac))
    ok, image = cap.read()
    cap.release()
    if not ok: raise RuntimeError(f"Could not read {video} at {frac}")
    cv2.imwrite(str(dst), image)
    return str(dst.relative_to(ROOT))


heat_fridge = frame("heat_breakfast.mp4", .15, "heat_fridge")
heat_microwave = frame("heat_breakfast.mp4", .45, "heat_microwave")
heat_table = frame("heat_breakfast.mp4", .75, "heat_table")
fruits = frame("collect_fruits.mp4", .48, "collect_fruits")
toys = frame("tidy_toys.mp4", .45, "tidy_toys")
books = frame("shelve_books.mp4", .16, "shelve_books")

# 1. Title
s = PPT.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY)
pic(s, "media/scenes/house_topdown.jpg", 6.20, .32, 6.72, 6.85, crop=True)
box(s, 6.08, .32, .09, 6.85, TEAL)
text(s, "ZenoBench", .65, 1.16, 5.12, .75, 36, WHITE, True)
text(s, "From generated worlds to household tasks", .65, 2.12, 5.17, 1.33, 25, WHITE, True)
text(s, "What we have · What it can do · What comes next", .66, 4.01, 5.20, .71, 16, RGBColor(194, 216, 234))
box(s, .65, 5.59, 4.98, .67, RGBColor(27, 57, 79), radius=True)
text(s, "Group meeting  |  October 2026", .88, 5.77, 4.48, .32, 14, WHITE, True)
text(s, "github.com/CMU-Agentic-Library/ZenoBench", .67, 6.83, 5.64, .25, 10.8, RGBColor(188, 211, 228))
note(s, "The talk is a high-level overview: upstream EmbodiedGen V2, what ZenoBench contains, task demos, and the proposed upper-level skill-graph interface.")

# 2. EmbodiedGen V2 in one slide
s = slide("EmbodiedGen V2 generates simulation-ready worlds", "Background", "The upstream system turns language or images into assets, rooms and task layouts.")
for i,(head,body,color) in enumerate([
    ("Assets", "3D objects with geometry, appearance and physical properties", TEAL),
    ("Scenes", "Navigable multi-room environments", BLUE),
    ("Task worlds", "Objects arranged for interactive robot tasks", AMBER),
]):
    x=.65+i*4.19
    box(s,x,2.09,3.73,2.49,PALE,LINE,radius=True)
    box(s,x,2.09,3.73,.13,color)
    text(s,head,x+.22,2.49,3.28,.47,23,INK,True)
    text(s,body,x+.22,3.17,3.27,1.03,15,MID)
    if i<2: text(s,"→",x+3.78,3.07,.28,.39,24,BLUE,True)
box(s,.65,5.12,12.07,.83,RGBColor(229,242,248),radius=True)
text(s,"ZenoBench uses these generated objects in a reusable house and turns them into runnable, measurable tasks.",.89,5.34,11.59,.40,17,NAVY,True)
text(s,"Source: HorizonRobotics / EmbodiedGen V2 README",.68,6.43,11.45,.22,10,MID)
note(s, "EmbodiedGen V2 is upstream. It documents text/image-to-3D assets, multi-room scenes and task-driven world composition: https://github.com/HorizonRobotics/EmbodiedGen/blob/master/README.md")

# 3. What is in ZenoBench
s = slide("ZenoBench at a glance", "What I built", "One reusable home, generated task objects, and a suite of executable household goals.")
things=[
    ("Assets", "EmbodiedGen V2 task objects prepared for physical interaction", "media/assets_gallery.jpg", TEAL),
    ("Rooms", "10-room Infinigen house with furniture and articulated storage", "media/scenes/house_topdown.jpg", BLUE),
    ("Tasks", "8 task definitions, each with a scene, goal and evaluation", heat_table, AMBER),
]
for i,(head,body,im,color) in enumerate(things):
    x=.61+i*4.19
    box(s,x,1.94,3.79,4.47,PALE,LINE,radius=True)
    pic(s,im,x+.10,2.04,3.59,2.28,crop=(i!=0))
    box(s,x+.10,4.40,3.59,.08,color)
    text(s,head,x+.22,4.73,3.34,.43,20,INK,True)
    text(s,body,x+.22,5.28,3.31,.81,13.6,MID)
note(s, "The house is an Infinigen layout, and the task objects are EmbodiedGen V2-generated. The repo prepares assets, room interactions, task scenes and scoring. This slide is intentionally a simple inventory, not an implementation pipeline.")

# 4. What it can do
s = slide("What can the robot do here?", "Capabilities", "Zeno Malo can execute connected household actions in the same environment.")
capabilities=[
    ("Move", "Travel between rooms and approach workspaces", "media/scenes/house_topdown.jpg", TEAL),
    ("Handle objects", "Pick, carry, place and sort different object shapes", fruits, BLUE),
    ("Interact", "Open/close storage and operate appliance controls", "tasks/heat_breakfast_combo/check/microwave_framed.png", AMBER),
    ("Complete tasks", "Combine actions toward a goal checked from simulator state", heat_table, GREEN),
]
for i,(head,body,im,color) in enumerate(capabilities):
    x=.61+(i%2)*6.27; y=1.97+(i//2)*2.31
    box(s,x,y,5.94,2.05,PALE,LINE,radius=True)
    pic(s,im,x+.09,y+.09,2.34,1.87,crop=True)
    box(s,x+2.55,y+.31,.07,1.46,color)
    text(s,head,x+2.80,y+.29,2.90,.40,18.2,INK,True)
    text(s,body,x+2.80,y+.81,2.91,.95,12.7,MID)
note(s, "Keep this high-level. These are things the system can demonstrate; do not explain trajectory controllers here. Recorded videos show whole scripted-policy runs and state-based evaluation.")

# Current atomic policies — deliberately above controller implementation detail.
s = slide("Atomic policies: what is callable today", "Current capabilities", "Concrete actions that can be combined under a task or a semantic contract.")
policies=[
    ("Navigate + move", "empty_navigate · carry_navigate\npick_while_moving · place_while_moving", TEAL),
    ("Posture + reach", "lower_torso · lean_forward\nright_tcp_move · gripper open/close", BLUE),
    ("Pick + place", "pick_top · pick_round_rim · pick_edge\nplace_surface · place_container", AMBER),
    ("Interact + operate", "open_handle · close_handle · click\nmicrowave_start · place_microwave", GREEN),
]
for i,(h,b,c) in enumerate(policies):
    x=.61+(i%2)*6.28; y=1.95+(i//2)*1.92
    box(s,x,y,5.93,1.69,PALE,LINE,radius=True)
    box(s,x,y,.09,1.69,c)
    text(s,h,x+.22,y+.20,5.43,.36,18.3,INK,True)
    text(s,b,x+.22,y+.72,5.36,.73,13.4,MID)
box(s,.64,6.04,12.00,.55,RGBColor(232,244,248),radius=True)
text(s,"60 OOP entries  ·  54 passed at least one Isaac Sim smoke test  ·  6 callable but unverified",.87,6.19,11.55,.29,14.6,NAVY,True)
text(s,"Unverified: floor-corner pick and five dual-arm / handover routes; verified results are scene-specific.",.71,6.69,11.87,.20,10.4,MID)
note(s, "Current repo catalog has 60 policy entries, 54 with at least one representative physical smoke run and six callable but unverified. Give examples by category, not the full list. The six unverified include floor-corner picking and five dual-arm or handover routes. An atomic policy is a graph-level action; its controller may include several joint motions.")

# 5. Suite categories
s = slide("Eight tasks across four household themes", "Task suite", "The same house supports different goals rather than a single scripted demonstration.")
cats=[
    ("Collect", "collect_fruits", "Gather objects into a shared container", TEAL),
    ("Organize", "tidy_toys\nshelve_books", "Put floor toys away; place books on shelves", BLUE),
    ("Prepare", "desk_prep\nbreakfast_setup", "Set up a workspace or a dining table", AMBER),
    ("Appliances", "heat_breakfast_preloaded\nheat_breakfast_combo\nheat_breakfast", "Operate fridge/microwave and serve food", GREEN),
]
for i,(head,names,body,color) in enumerate(cats):
    x=.59+i*3.14
    box(s,x,2.02,2.76,3.93,PALE,LINE,radius=True)
    box(s,x,2.02,2.76,.13,color)
    text(s,head,x+.19,2.42,2.40,.39,20,INK,True)
    text(s,names,x+.19,3.12,2.42,1.17,13.5,color,True)
    text(s,body,x+.19,4.72,2.39,.94,13,MID)
text(s,"Representative rollouts and result JSONs are included in the repository.",.66,6.41,11.98,.32,14,NAVY,True)
note(s, "All eight names come from task_specs. There are five organization/preparation tasks and three appliance tasks. The README includes representative runs; the task list is not a claim of robust success across all seeds.")

# 6. Task examples
s = slide("Task showcase: three different behaviors", "Demo", "Use the recorded videos to show the system acting in the house.")
stories=[
    (fruits,"Collect fruits","Cross rooms; put two fruits into one container",TEAL,"collect_fruits"),
    (toys,"Tidy toys","Pick floor objects and store them",BLUE,"tidy_toys"),
    (books,"Shelve books","Handle flat books and place them on shelves",AMBER,"shelve_books"),
]
for i,(im,h,b,color,video) in enumerate(stories):
    x=.60+i*4.18
    box(s,x,1.99,3.80,4.40,PALE,LINE,radius=True)
    preview=pic(s,im,x+.09,2.08,3.62,2.29)
    preview.click_action.hyperlink.address = f"https://github.com/CMU-Agentic-Library/ZenoBench/blob/main/media/tasks/{video}.mp4"
    box(s,x+.09,4.46,3.62,.08,color)
    text(s,h,x+.23,4.79,3.35,.41,19,INK,True)
    text(s,b,x+.23,5.36,3.34,.73,13.4,MID)
note(s, "If time is limited, play only one clip. These stills come from the recorded collect_fruits, tidy_toys and shelve_books videos. Video paths are listed in SPEAKER_NOTES.md.")

# 7. Hero task
s = slide("One longer example: fridge → microwave → table", "Task demo", "A single goal ties together object handling, appliance use and movement through the home.")
for i,(im,head) in enumerate([(heat_fridge,"Take from fridge"),(heat_microwave,"Heat in microwave"),(heat_table,"Serve at table")]):
    x=.59+i*4.18
    box(s,x,1.99,3.80,3.50,PALE,LINE,radius=True)
    preview=pic(s,im,x+.09,2.08,3.62,2.31)
    preview.click_action.hyperlink.address = "https://github.com/CMU-Agentic-Library/ZenoBench/blob/main/media/tasks/heat_breakfast.mp4"
    text(s,head,x+.21,4.68,3.40,.40,17.5,INK,True)
    if i<2: text(s,"→",x+3.86,3.05,.28,.41,24,BLUE,True)
box(s,.66,5.83,12.01,.58,RGBColor(231,246,238),radius=True)
text(s,"Recorded run: all task goals met; oatmeal served warm, doors closed, no object dropped.",.90,5.99,11.52,.31,15,GREEN,True)
text(s,"One representative rollout; not a multi-seed benchmark result.",.69,6.60,11.92,.25,10.8,MID)
note(s, "The heat_breakfast seed-0 rollout achieved 100% of its four goals. The repo video shows the full sequence. Heating is modeled at task level, not with PhysX thermal simulation. Avoid quoting this as general success rate.")

# Representative whole-task status, without implying multi-seed robustness.
s = slide("What whole tasks run today?", "Task results", "Eight task definitions; results shown are representative recorded rollouts.")
records=[
    ("collect_fruits", "Two fruits into one container", "SUCCESS", GREEN),
    ("tidy_toys", "Floor toys into toy box / basket", "SUCCESS", GREEN),
    ("shelve_books", "Flat books onto a bookcase", "SUCCESS", GREEN),
    ("desk_prep", "Notebook, pen and mug on desk", "SUCCESS", GREEN),
    ("breakfast_setup", "Dining table place setting", "67%", AMBER),
    ("heat_breakfast_preloaded", "Microwave heating from preloaded food", "SUCCESS", GREEN),
    ("heat_breakfast_combo", "Microwave cycle plus chilled milk", "SUCCESS", GREEN),
    ("heat_breakfast", "Fridge → microwave → dining table", "SUCCESS", GREEN),
]
for i,(name,desc,status,color) in enumerate(records):
    col=i//4; row=i%4
    x=.61+col*6.27; y=1.93+row*1.03
    box(s,x,y,5.93,.87,PALE,LINE,radius=True)
    box(s,x,y,.075,.87,color)
    text(s,name,x+.19,y+.14,3.48,.28,13.6,INK,True)
    text(s,desc,x+.19,y+.47,4.65,.25,11.2,MID)
    box(s,x+4.83,y+.18,.89,.36,RGBColor(230,245,236) if status=="SUCCESS" else RGBColor(252,239,216),radius=True)
    text(s,status,x+4.90,y+.28,.76,.20,9.4,color,True,align=PP_ALIGN.CENTER)
box(s,.66,6.23,12.00,.51,RGBColor(236,246,250),radius=True)
text(s,"7 successful representative runs; breakfast_setup remains partial after a mug drop.",.90,6.38,11.49,.27,14.5,NAVY,True)
note(s, "These are recorded examples, not a statistical benchmark result. The repo has six successful videos in the main task table, a separate successful preloaded microwave smoke run, and one partial breakfast_setup run at 67% after a mug drop. The full heat_breakfast task has a seed-0 successful rollout.")

# 8. Why useful
s = slide("Why this matters for our setting", "Research value", "These tasks give an upper-level planner a meaningful environment to reason, act and recover in.")
for i,(head,body,color) in enumerate([
    ("Plan", "Choose a sequence of skills for a household goal", TEAL),
    ("Ground", "Bind each skill to objects, rooms and executable actions", BLUE),
    ("Verify & replan", "Observe the outcome, detect failure and revise the next step", AMBER),
]):
    x=.61+i*4.18
    box(s,x,2.13,3.79,3.35,PALE,LINE,radius=True)
    box(s,x,2.13,3.79,.13,color)
    text(s,f"0{i+1}",x+.22,2.48,.57,.40,19,color,True)
    text(s,head,x+.22,3.12,3.30,.49,22,INK,True)
    text(s,body,x+.22,3.80,3.31,1.18,15,MID)
text(s,"ZenoBench is the execution and evaluation substrate for the upper-level Skill Graph design.",.68,6.02,11.98,.53,18,NAVY,True)
note(s, "This is the transition to the user's two architecture figures. The existing task suite tests the consequences of upper-level decisions. The upper layer will reason over a public skill library and inspect verified outcomes.")

# Upper-level handoff: the two supplied architecture figures are reconstructed
# at full resolution in presentations/figures/ and inserted as standalone slides.
for figure, note_text in [
    ("presentations/figures/upper_figure_1_skill_graph.png",
     "Handoff to the upper-layer presenter: Figure 1 shows one Skill Subgraph per subgoal, validation, execution, verification and replanning."),
    ("presentations/figures/upper_figure_2_skill_library.png",
     "Handoff to the upper-layer presenter: Figure 2 shows public SkillNodes, per-skill Contracts, the invocation adapter and heterogeneous low-level policies."),
]:
    s = PPT.slides.add_slide(BLANK)
    box(s,0,0,13.333,7.5,WHITE)
    pic(s,figure,0,0,13.333,7.5,crop=False)
    note(s,note_text)

PPT.core_properties.title = "ZenoBench: Generated Worlds to Household Tasks"
PPT.core_properties.subject = "Group meeting overview, current atomic policies, task demos and upper-level handoff"
PPT.core_properties.keywords = "ZenoBench; EmbodiedGen V2; atomic policy; Skill Graph"
PPT.save(DEST)
print(DEST)
print(f"slides={len(PPT.slides)}")
